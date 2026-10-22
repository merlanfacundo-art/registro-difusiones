"""
Genera api/_data/historico.ts a partir de las solapas mensuales viejas
(Enero-Julio) de "Organización 2026", exportada como .xlsx.

Uso (Windows):
    pip install openpyxl
    python scripts\\generar_historico.py "C:\\ruta\\Organizacion 2026.xlsx"

Solo hace falta volver a correrlo si se corrige algún dato de esas
solapas viejas: no se cargan más, así que la base es una foto fija.

Criterios (los mismos de la planilla "Estadísticas combinadas"):
- Una fila por compañerx por actividad. Celda vacía = "Sin Respuesta".
- Se excluyen columnas de actividad sin ninguna respuesta cargada
  (placeholders de Julio y "2da capacitación ASIS", que vive en el sistema nuevo).
- ✅ / ❌ se toman del comentario de la actividad.
- El día de semana se infiere del nombre de la actividad, validando el mes;
  si no se puede con confianza, queda en null.
"""
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

MESES = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio']
ETIQUETAS_PIE = {'Confirmadxs', 'Ausentes', 'A confirmar', 'Sin Respuesta', 'N/A'}
DIAS = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']
# Códigos compactos de respuesta (el frontend usa el mismo orden)
CODIGOS = {'Si': 0, 'No': 1, 'A confirmar': 2, 'Sin Respuesta': 3, 'N/A': 4}
ANIO = 2026


def inferir_dia(nombre: str, mes_num: int):
    # Si el nombre menciona más de un día (actividad de varias jornadas),
    # no se asigna ninguno: elegir uno sería inventar.
    if len(re.findall(r'(Lunes|Martes|Mi[eé]rcoles|Jueves|Viernes|S[aá]bado|Sab[aá]do|Domingo)', nombre, re.IGNORECASE)) > 1:
        return None
    m = re.search(r'-\s*(Lunes|Martes|Mi[eé]rcoles|Jueves|Viernes|S[aá]bado|Sab[aá]do|Domingo)\s+(\d{1,2})',
                  nombre, re.IGNORECASE)
    if m:
        try:
            return DIAS[datetime.date(ANIO, mes_num, int(m.group(2))).weekday()]
        except ValueError:
            return None
    m = re.search(r'-\s*(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?', nombre)
    if m and int(m.group(2)) == mes_num:
        try:
            return DIAS[datetime.date(ANIO, mes_num, int(m.group(1))).weekday()]
        except ValueError:
            return None
    return None


def main(ruta_xlsx: str):
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    compas, areas, actividades, registros = [], [], [], []
    idx_compa_map, idx_area_map = {}, {}

    def indice(lista, mapa, valor):
        if valor not in mapa:
            mapa[valor] = len(lista)
            lista.append(valor)
        return mapa[valor]

    for mes_num, mes in enumerate(MESES, start=1):
        ws = wb[mes]
        header = [c.value for c in ws[1]]
        idx_compa = header.index('Compa/Actividad')
        idx_area = 0 if idx_compa >= 1 and str(header[0]).strip() in ('Area', 'Área') else None
        filas = [r for r in ws.iter_rows(min_row=2, values_only=True)
                 if r[idx_compa] is not None and r[idx_compa] not in ETIQUETAS_PIE]

        col = idx_compa + 1
        while col < len(header) and header[col] is not None and header[col] not in ETIQUETAS_PIE:
            nombre = str(header[col]).strip()
            if any(f[col] is not None for f in filas):
                ai = len(actividades)
                actividades.append({'n': nombre, 'm': mes, 'd': inferir_dia(nombre, mes_num)})
                for f in filas:
                    resp = str(f[col]).strip() if f[col] is not None else 'Sin Respuesta'
                    if resp not in CODIGOS:
                        raise ValueError(f'Respuesta no reconocida en {mes} / {nombre}: {resp!r}')
                    coment = str(f[col + 1]) if col + 1 < len(f) and f[col + 1] is not None else ''
                    flags = (1 if '✅' in coment else 0) | (2 if '❌' in coment else 0)
                    area = f[idx_area] if idx_area is not None else None
                    registros.append([
                        indice(compas, idx_compa_map, str(f[idx_compa]).strip()),
                        ai,
                        indice(areas, idx_area_map, str(area).strip()) if area else -1,
                        CODIGOS[resp],
                        flags,
                    ])
            col += 2

    data = {'compas': compas, 'areas': areas, 'actividades': actividades, 'registros': registros}
    salida = Path(__file__).resolve().parent.parent / 'api' / '_data' / 'historico.ts'
    salida.parent.mkdir(parents=True, exist_ok=True)
    salida.write_text(
        '// Generado por scripts/generar_historico.py — no editar a mano.\n'
        '// registros: [compañerx, actividad, área (-1 = sin dato), respuesta, flags (1=✅, 2=❌)]\n'
        '// respuesta: 0=Si 1=No 2=A confirmar 3=Sin Respuesta 4=N/A\n'
        'export const HISTORICO = '
        + json.dumps(data, ensure_ascii=False, separators=(',', ':'))
        + ';\n',
        encoding='utf-8',
    )
    print(f'{len(actividades)} actividades, {len(registros)} registros -> {salida}')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit('Uso: python scripts/generar_historico.py "ruta/Organizacion 2026.xlsx"')
    main(sys.argv[1])
